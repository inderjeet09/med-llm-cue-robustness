"""Analyse the 80-item validity review (pre-registered in docs/preregistration.md: invalid items are excluded and a validity
rate reported; templates are never edited).

Input  data/processed/validity_review/review_sheet_completed.xlsx (+ key.csv: row -> variant id, family)
Output results/review/review_summary.json and review_summary.md
       Sensitivity: H1-type differences on held-out cues recomputed with the source questions of items judged invalid
       (Q1 = N) or with a wrong key (Q2 = N) removed from the test set.
Usage: python scripts/analyze_review.py
"""
from __future__ import annotations

import csv
import json
import statistics
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from src.data.bias import FAMILY_OF, NEW_BIASES  # noqa: E402
from src.utils.provenance import provenance  # noqa: E402

D = ROOT / "data/processed/validity_review"


def main() -> None:
    from openpyxl import load_workbook
    ws = load_workbook(D / "review_sheet_completed.xlsx", data_only=True)["Review"]
    rows = [r for r in list(ws.iter_rows(values_only=True))[1:] if r and r[0] is not None]
    key = {int(k["Row"]): k for k in csv.DictReader((D / "key.csv").open())}
    items = []
    for r in rows:
        k = key[int(r[0])]
        bias = k["variant"].split(":")[1] if ":" in k["variant"] else None
        items.append({"row": int(r[0]), "variant": k["variant"], "variant_id": k["variant_id"],
                      "src_id": k["variant_id"].split("|")[0], "kind": "nota" if k["variant"] == "nota" else "hint",
                      "bias": bias, "family": FAMILY_OF.get(bias) if bias else None,
                      "new_family": bias in NEW_BIASES if bias else None,
                      "q1": str(r[9]).strip().upper() if r[9] is not None else None,
                      "q2": str(r[10]).strip().upper() if r[10] is not None else None,
                      "q3": int(r[11]) if r[11] not in (None, "") else None,
                      "q4": str(r[12]).strip().upper() if r[12] not in (None, "") else None,
                      "comment": r[13]})
    hint = [i for i in items if i["kind"] == "hint"]
    nota = [i for i in items if i["kind"] == "nota"]
    q3 = [i["q3"] for i in hint if i["q3"] is not None]
    fam = defaultdict(list)
    for i in hint:
        if i["q3"] is not None:
            fam[i["family"]].append(i["q3"])
    flagged = [i for i in items if i["q1"] == "N" or i["q2"] == "N"]
    res = {
        "n_items": len(items), "n_hint": len(hint), "n_nota": len(nota),
        "q1_valid": sum(i["q1"] == "Y" for i in items), "q2_key_best": sum(i["q2"] == "Y" for i in items),
        "q2_key_best_hint": sum(i["q2"] == "Y" for i in hint), "q2_key_best_nota": sum(i["q2"] == "Y" for i in nota),
        "q4_nota_clearly_correct": sum(i["q4"] == "Y" for i in nota),
        "q3_plausibility_mean": statistics.mean(q3), "q3_plausibility_median": statistics.median(q3),
        "q3_distribution": {s: q3.count(s) for s in range(1, 6)},
        "q3_share_ge3": 100 * sum(v >= 3 for v in q3) / len(q3),
        "q3_mean_by_family": {f: statistics.mean(v) for f, v in sorted(fam.items())},
        "q3_mean_biasmedqa": statistics.mean([i["q3"] for i in hint if i["q3"] is not None and not i["new_family"]]),
        "q3_mean_new": statistics.mean([i["q3"] for i in hint if i["q3"] is not None and i["new_family"]]),
        "flagged_items": [{k: i[k] for k in ("row", "variant", "src_id", "q1", "q2", "q4", "comment")} for i in flagged],
        "n_comments": sum(bool(i["comment"]) for i in items),
    }

    # sensitivity: remove the source questions of flagged items from the test set, recompute held-out-cue differences
    from analyze_main import collect
    from src.evaluation.stats import two_level_bootstrap
    drop = {i["src_id"] for i in flagged}
    sens = {}
    for m in ["Qwen__Qwen2.5-7B-Instruct", "meta-llama__Llama-3.1-8B-Instruct", "m42-health__Llama3-Med42-8B"]:
        A = collect(ROOT / "results/main" / m, "test")
        q = A["cift"][0]["_qids"]
        keep = np.array([qq not in drop for qq in q])
        out = {"n_questions_kept": int(keep.sum())}
        for other in ("base", "sft_clean", "bct"):
            r = two_level_bootstrap(np.stack([x["heldout_cues"][keep] for x in A["cift"]]),
                                    np.stack([x["heldout_cues"][keep] for x in A[other]]), seed=0)
            out[f"cift_minus_{other}_heldout"] = {"diff_pp": 100 * r["diff"], "ci95_pp": [100 * v for v in r["ci95"]]}
        sens[m] = out
    res["sensitivity_flagged_removed"] = sens
    res["provenance"] = provenance({}, __file__)
    out_dir = ROOT / "results/review"; out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "review_summary.json").write_text(json.dumps(res, indent=1, default=str))

    L = ["# Item validity review (80 test items)", "",
         f"- Q1 valid and answerable: {res['q1_valid']}/{res['n_items']}",
         f"- Q2 answer key still the single best answer: {res['q2_key_best']}/{res['n_items']} "
         f"(hint items {res['q2_key_best_hint']}/{res['n_hint']}, NOTA items {res['q2_key_best_nota']}/{res['n_nota']})",
         f"- Q4 NOTA clearly correct: {res['q4_nota_clearly_correct']}/{res['n_nota']}",
         f"- Q3 hint plausibility (1-5): mean {res['q3_plausibility_mean']:.2f}, median {res['q3_plausibility_median']}, "
         f"distribution {res['q3_distribution']}, rated >= 3: {res['q3_share_ge3']:.0f} %",
         f"- Q3 mean, BiasMedQA templates {res['q3_mean_biasmedqa']:.2f} vs new templates {res['q3_mean_new']:.2f}; by family: "
         + ", ".join(f"{f} {v:.2f}" for f, v in res["q3_mean_by_family"].items()),
         "", "## Flagged items (Q1 or Q2 = N)"] + [f"- row {i['row']} {i['variant']} ({i['src_id']}): Q1 {i['q1']}, Q2 {i['q2']}, "
                                                  f"Q4 {i['q4']}; {i['comment']}" for i in res["flagged_items"]] + \
        ["", "## Sensitivity: held-out-cue differences with flagged source questions removed"] + \
        [f"- {m.split('__')[1]} (n={v['n_questions_kept']}): " + "; ".join(
            f"CIFT−{k.split('_minus_')[1].replace('_heldout', '')} {d['diff_pp']:+.1f} [{d['ci95_pp'][0]:.1f}, {d['ci95_pp'][1]:.1f}]"
            for k, d in v.items() if k.startswith("cift")) for m, v in sens.items()]
    (out_dir / "review_summary.md").write_text("\n".join(L) + "\n")
    print("\n".join(L))


if __name__ == "__main__":
    main()
