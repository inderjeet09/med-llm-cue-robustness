"""Phase 2: random sample of TEST-split variants for the item validity review (a medically trained reviewer, D028).

80 rows, each from a different source question: 4 misleading-cue items per template (15 templates = 60) and 20
NOTA items. The reviewer judges medical validity only; answers are not used to tune anything. Frozen templates are
not changed in response to this review: invalid items are reported (validity rate) and can be excluded, never
edited (pre-registered in docs/preregistration.md).

Output (gitignored, contains test questions): data/processed/validity_review/review_sheet.xlsx (+ .csv) and
key.csv (row -> variant id). Instructions: results/review/validity_review_instructions.md.
Usage: python scripts/make_review_sheet.py
"""
from __future__ import annotations

import csv
import json
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from src.data.bias import TEMPLATES  # noqa: E402

SEED = 2026
PER_TEMPLATE, N_NOTA = 4, 20
COLS = ["Row", "Item type", "Question (exactly as the model sees it)", "A", "B", "C", "D", "Answer key",
        "Hint points to", "Q1 Valid & answerable? (Y/N)", "Q2 Answer key still the single best answer? (Y/N)",
        "Q3 Hint sentence clinically plausible? (1-5, hint rows)", "Q4 NOTA now clearly correct? (Y/N, NOTA rows)",
        "Comments"]


INSTRUCTIONS = [
    "Review of 80 medical exam questions (about 1-2 hours)",
    "",
    "Why: we test whether AI models can be misled when ONE sentence is added to a USMLE-style exam question (MedQA, a public "
    "dataset with no patient data), or when the correct option is replaced by \"None of the other answers\". These versions "
    "were generated automatically. We need a medical reviewer to tell us how often the edits make a question medically "
    "invalid. Your answers are used only to report a validity rate and to exclude invalid items; they never change the AI.",
    "",
    "What to do: open the \"Review\" tab. Each row is one question with its options A-D and the answer key. Fill in the four "
    "yellow columns (Q1-Q4) for every row, plus Comments if something is wrong. Work on your own; please don't look answers "
    "up beyond normal medical knowledge, and don't discuss the items with others.",
    "",
    "Two kinds of rows (column \"Item type\"):",
    "  - Misleading hint: one sentence was added that points to a WRONG option (column \"Hint points to\"). The answer key is "
    "the original exam answer.",
    "  - Correct answer replaced by NOTA: the correct option's text was replaced by \"None of the other answers\". The answer "
    "key is that option.",
    "",
    "Q1  Is the question still medically valid and answerable as written?  (Y / N)",
    "Q2  Is the answer key still the single best answer? Judge the medicine, ignore what the hint suggests.  (Y / N)",
    "Q3  Hint rows only: does the added sentence read as something that could plausibly be said in a clinical setting?  "
    "(1 = absurd ... 5 = entirely plausible)",
    "Q4  NOTA rows only: is \"None of the other answers\" now clearly correct, i.e. none of the remaining options is "
    "acceptable?  (Y / N)",
    "Comments: anything wrong with the item (free text).",
    "",
    "When done: save the file and send it back. Thank you!",
]


def main() -> None:
    rows = [json.loads(l) for l in open(ROOT / "data/processed/medqa_us4_variants/test.jsonl")]
    rng = random.Random(SEED)
    src_ids = sorted({r["src_id"] for r in rows})
    rng.shuffle(src_ids)
    by = {(r["src_id"], r["variant"]): r for r in rows}
    picks, it = [], iter(src_ids)
    for bias in TEMPLATES:
        for _ in range(PER_TEMPLATE):
            picks.append(by[(next(it), f"mis:{bias}")])
    for _ in range(N_NOTA):
        picks.append(by[(next(it), "nota")])
    rng.shuffle(picks)  # mix item types so the reviewer does not see blocks of one template

    out = ROOT / "data/processed/validity_review"
    out.mkdir(parents=True, exist_ok=True)
    table, key = [], []
    for i, r in enumerate(picks, 1):
        kind = "Misleading hint" if r["variant"].startswith("mis:") else "Correct answer replaced by NOTA"
        table.append([i, kind, r["question"], *(r["options"][k] for k in "ABCD"), r["answer_idx"],
                      r.get("cue_target") or "", "", "", "", "", ""])
        key.append({"Row": i, "variant_id": r["id"], "variant": r["variant"], "family": r.get("family", "")})
    with (out / "review_sheet.csv").open("w", newline="") as f:
        csv.writer(f).writerows([COLS, *table])
    with (out / "key.csv").open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(key[0]))
        w.writeheader(); w.writerows(key)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ins = wb.active; ins.title = "Instructions"   # first tab: everything the reviewer needs, so only one file is sent
    for line in INSTRUCTIONS:
        ins.append([line])
    ins.column_dimensions["A"].width = 120
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ins["A1"].font = Font(bold=True, size=14)
    ws = wb.create_sheet("Review")
    ws.append(COLS)
    for row in table:
        ws.append(row)
    widths = [5, 16, 70, 22, 22, 22, 22, 8, 8, 12, 14, 14, 14, 30]
    for col, w in zip("ABCDEFGHIJKLMN", widths):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor="E2F0EF")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=10, max_col=13):
        for c in row:
            c.fill = PatternFill("solid", fgColor="FFF4C2")
    ws.freeze_panes = "C2"
    wb.save(out / "review_sheet.xlsx")
    print(f"{len(table)} rows -> {out.relative_to(ROOT)}/review_sheet.xlsx (+ .csv, key.csv)")


if __name__ == "__main__":
    main()
